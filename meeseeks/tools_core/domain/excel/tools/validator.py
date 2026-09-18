"""
Validator - Post-update verification and audit logging
"""

import os
import json
import hashlib
import logging
import openpyxl
from pathlib import Path
from datetime import datetime
from typing import Optional
from dataclasses import asdict

from ..core.data_classes import UpdatePlan, ExecutionResult, ValidationResult, AuditLog

logger = logging.getLogger(__name__)


class Validator:
    """
    Post-update verification and audit logging.
    """
    
    def __init__(self, audit_dir: Optional[Path] = None):
        # Default to logs/audit_logs relative to project root (not CWD!)
        default_dir = Path(__file__).parent.parent.parent.parent.parent / "logs" / "audit_logs"
        self.audit_dir = audit_dir or default_dir
        self.audit_dir.mkdir(exist_ok=True)
    
    def validate(self, original_path: Path, updated_path: Path, 
                 plan: UpdatePlan) -> ValidationResult:
        """Validate that updates were applied correctly"""
        logger.info("✅ Validating updates...")
        
        issues = []
        formula_errors = []
        checks_passed = 0
        checks_failed = 0
        
        try:
            wb_original = openpyxl.load_workbook(original_path, data_only=True)
            wb_updated = openpyxl.load_workbook(updated_path, data_only=True)
            
            for target in plan.targets:
                sheet = wb_updated[target.location.sheet]
                cell = sheet[target.location.cell]
                
                if cell.value == target.new_value:
                    checks_passed += 1
                else:
                    checks_failed += 1
                    issues.append(f"Value mismatch at {target.location.sheet}!{target.location.cell}")
            
            wb_updated_formulas = openpyxl.load_workbook(updated_path, data_only=True)
            for sheet_name in wb_updated_formulas.sheetnames:
                sheet = wb_updated_formulas[sheet_name]
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value in ['#VALUE!', '#REF!', '#NAME?', '#DIV/0!', '#NULL!', '#N/A']:
                            formula_errors.append(f"{sheet_name}!{cell.coordinate}: {cell.value}")
            
            wb_original.close()
            wb_updated.close()
            wb_updated_formulas.close()
            
        except Exception as e:
            issues.append(f"Validation error: {e}")
            checks_failed += 1
        
        valid = checks_failed == 0 and len(formula_errors) == 0
        
        logger.info(f"  ✓ Checks passed: {checks_passed}")
        if checks_failed > 0:
            logger.warning(f"  ✗ Checks failed: {checks_failed}")
        if formula_errors:
            logger.warning(f"  ⚠ Formula errors: {len(formula_errors)}")
        
        return ValidationResult(
            valid=valid,
            checks_passed=checks_passed,
            checks_failed=checks_failed,
            issues=issues,
            formula_errors=formula_errors[:20],
            metric_changes={}
        )
    
    def create_audit_log(self, plan: UpdatePlan, result: ExecutionResult, 
                         validation: ValidationResult) -> AuditLog:
        """Create immutable audit record"""
        log_id = f"audit_{plan.id}"
        timestamp = datetime.now().isoformat()
        
        content = json.dumps({
            'plan_id': plan.id,
            'result': asdict(result),
            'validation': asdict(validation),
            'timestamp': timestamp
        }, sort_keys=True)
        checksum = hashlib.sha256(content.encode()).hexdigest()
        
        audit_log = AuditLog(
            id=log_id,
            timestamp=timestamp,
            user=os.getenv('USER', 'unknown'),
            action='update',
            plan=plan,
            result=result,
            validation=validation,
            checksum=checksum
        )
        
        log_path = self.audit_dir / f"{log_id}.json"
        with open(log_path, 'w') as f:
            json.dump(self._audit_to_dict(audit_log), f, indent=2, default=str)
        
        logger.info(f"📝 Audit log saved: {log_path.name}")
        
        return audit_log
    
    def _audit_to_dict(self, audit: AuditLog) -> dict:
        """Convert audit log to dictionary"""
        return {
            'id': audit.id,
            'timestamp': audit.timestamp,
            'user': audit.user,
            'action': audit.action,
            'plan': {
                'id': audit.plan.id,
                'source_file': audit.plan.source_file,
                'excel_file': audit.plan.excel_file,
                'target_count': len(audit.plan.targets),
                'reasoning_trace': audit.plan.reasoning_trace
            },
            'result': asdict(audit.result),
            'validation': asdict(audit.validation),
            'checksum': audit.checksum
        }

